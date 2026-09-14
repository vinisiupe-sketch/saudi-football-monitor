"""O download precisa ser um XLSX filtrável, legível e seguro."""
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import sys
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from glossario_excel import criar_excel_glossario


class ExportacaoExcel(unittest.TestCase):
    def test_planilha_guarda_ids_nomes_datas_e_filtro(self):
        conteudo = criar_excel_glossario([{
            "id": 9, "nome_principal": "=nome externo", "clube": "Al Hilal",
            "posicao": "Forward", "camisa": "10", "spl_id": "123",
            "af_id": 456, "tm_id": 789, "status": "em_preparo",
            "revisado": False, "nascimento": "2001-02-03",
            "atualizado_em": "2026-09-14T10:30:00+00:00",
            "nomes": [
                {"fonte": "spl", "idioma": "lat", "nome": "Player Name"},
                {"fonte": "pdf", "idioma": "lat", "nome": "PLAYER NAME"},
            ],
        }])
        wb = load_workbook(BytesIO(conteudo))
        ws = wb["Glossário"]
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws["A2"].value, 9)
        self.assertEqual(ws["B2"].value, "'=nome externo")
        self.assertEqual(ws["N2"].value, "PLAYER NAME")
        self.assertIsInstance(ws["R2"].value, (date, datetime))
        self.assertIsInstance(ws["S2"].value, datetime)
        self.assertIn("GlossarioJogadores", ws.tables)


if __name__ == "__main__":
    unittest.main()
