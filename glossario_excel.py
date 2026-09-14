"""Exportação simples e legível do laboratório para Excel."""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


COLUNAS = [
    "ID interno", "Jogador", "Clube", "Posição", "Camisa", "ID SPL",
    "SPL em inglês", "SPL em árabe", "ID API-Football", "API-Football",
    "ID Transfermarkt", "Transfermarkt", "TM em árabe", "PDF",
    "Notícias árabes", "Situação", "Revisado", "Nascimento", "Atualizado em",
]


def _nomes(jogador, fonte, idioma=""):
    vistos = []
    for nome in jogador.get("nomes") or []:
        if nome.get("fonte") != fonte or (idioma and nome.get("idioma") != idioma):
            continue
        texto = " ".join((nome.get("nome") or "").split())
        if texto and texto not in vistos:
            vistos.append(texto)
    return " | ".join(vistos)


def _texto_seguro(valor):
    if not isinstance(valor, str):
        return valor
    # Evita que uma grafia externa seja interpretada como fórmula pelo Excel.
    return "'" + valor if valor.startswith(("=", "+", "-", "@")) else valor


def _data_excel(valor, apenas_data=False):
    """Conserva datas como datas reais, para o Excel ordenar corretamente."""
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        convertido = valor
    elif isinstance(valor, date):
        return valor
    else:
        try:
            convertido = datetime.fromisoformat(str(valor).replace("Z", "+00:00"))
        except ValueError:
            try:
                return date.fromisoformat(str(valor))
            except ValueError:
                return _texto_seguro(str(valor))
    if apenas_data:
        return convertido.date()
    # O formato XLSX não aceita timezone dentro da célula.
    return convertido.replace(tzinfo=None)


def criar_excel_glossario(jogadores: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Glossário"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.append(COLUNAS)
    for jogador in jogadores:
        linha = [
            jogador.get("id"), jogador.get("nome_principal"), jogador.get("clube"),
            jogador.get("posicao"), jogador.get("camisa"), jogador.get("spl_id"),
            _nomes(jogador, "spl", "lat"), _nomes(jogador, "spl", "ar"),
            jogador.get("af_id"), _nomes(jogador, "api_football"),
            jogador.get("tm_id"), _nomes(jogador, "transfermarkt", "lat"),
            _nomes(jogador, "transfermarkt", "ar"), _nomes(jogador, "pdf"),
            _nomes(jogador, "noticia", "ar"),
            (jogador.get("status") or "").replace("_", " "),
            "Sim" if jogador.get("revisado") else "Não",
            _data_excel(jogador.get("nascimento"), apenas_data=True),
            _data_excel(jogador.get("atualizado_em")),
        ]
        ws.append([_texto_seguro(v) for v in linha])

    verde = "B6FF00"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17171E")
        cell.font = Font(color=verde, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    larguras = [12, 32, 20, 16, 10, 36, 32, 30, 17, 28, 18, 28, 30, 34, 34, 16, 12, 14, 22]
    for indice, largura in enumerate(larguras, 1):
        ws.column_dimensions[ws.cell(1, indice).column_letter].width = largura
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for coluna in (18, 19):
        for cell in ws.iter_cols(min_col=coluna, max_col=coluna, min_row=2):
            for item in cell:
                if hasattr(item.value, "year"):
                    item.number_format = "dd/mm/yyyy" if coluna == 18 else "dd/mm/yyyy hh:mm"
    ultima = max(2, ws.max_row)
    if ws.max_row == 1:
        ws.append([""] * len(COLUNAS))
    tabela = Table(displayName="GlossarioJogadores", ref=f"A1:S{ultima}")
    tabela.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False)
    ws.add_table(tabela)
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
